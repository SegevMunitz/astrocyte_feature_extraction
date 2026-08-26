#!/usr/bin/env python3
"""Build revised BMP50 morph → RNA PowerPoint."""

from __future__ import annotations

import json
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "slides" / "assets"
MORPH = ASSETS / "morph"
RNA = ASSETS / "rnaseq"
GEN = ASSETS / "generated"
OUT = ROOT / "slides" / "BMP50_morph_to_RNA.pptx"

TIME_ORDER = ("ctrl", "4h", "24h", "72h", "7d")
TIME_COLORS = {
    "ctrl": "#4c4c4c",
    "4H": "#1f77b4",
    "4h": "#1f77b4",
    "24H": "#ff7f0e",
    "24h": "#ff7f0e",
    "72H": "#2ca02c",
    "72h": "#2ca02c",
    "7D": "#d62728",
    "7d": "#d62728",
}
CORE = ("BMP_Id", "cell_cycle", "reactive_GFAP", "cytoskeleton", "OPC_pdgfra", "late_7d")
SUPP = ("MHC_class_I", "Stat3_reactive", "glutamate_ion", "ECM_remodeling")


def set_run(p, text: str, size: int = 18, bold: bool = False, color=(0x22, 0x22, 0x22)):
    p.clear()
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = "Calibri"


def add_title_slide(prs, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.1), Inches(12), Inches(1.4))
    set_run(box.text_frame.paragraphs[0], title, size=30, bold=True)
    box2 = slide.shapes.add_textbox(Inches(0.7), Inches(3.6), Inches(12), Inches(2.2))
    tf = box2.text_frame
    tf.word_wrap = True
    for i, line in enumerate(subtitle.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_run(p, line, size=16, color=(0x55, 0x55, 0x55))
        p.space_after = Pt(6)


def add_bullets(prs, title: str, bullets: list[str], size: int = 16) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = slide.shapes.add_textbox(Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.65))
    set_run(t.text_frame.paragraphs[0], title, size=24, bold=True)
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.05), Inches(12), Inches(6.1))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_run(p, f"•  {line}", size=size)
        p.space_after = Pt(8)


def add_full_image(prs, title: str, image: Path, caption: str = "", title_size: int = 20) -> None:
    """Image almost full-slide under a short title."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = slide.shapes.add_textbox(Inches(0.35), Inches(0.12), Inches(12.6), Inches(0.45))
    set_run(t.text_frame.paragraphs[0], title, size=title_size, bold=True)
    if image.is_file():
        # Widescreen content area ~12.6 x 6.4
        slide.shapes.add_picture(str(image), Inches(0.35), Inches(0.6), width=Inches(12.6))
    if caption:
        c = slide.shapes.add_textbox(Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.35))
        set_run(c.text_frame.paragraphs[0], caption, size=11, color=(0x66, 0x66, 0x66))


def add_image_with_notes(prs, title: str, image: Path, notes: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    t = slide.shapes.add_textbox(Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.45))
    set_run(t.text_frame.paragraphs[0], title, size=20, bold=True)
    if image.is_file():
        slide.shapes.add_picture(str(image), Inches(0.35), Inches(0.65), width=Inches(8.4))
    notes_box = slide.shapes.add_textbox(Inches(8.95), Inches(0.7), Inches(4.0), Inches(6.3))
    tf = notes_box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(notes):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        set_run(p, f"•  {line}", size=13)
        p.space_after = Pt(8)


def make_volcano(csv_path: Path, title: str, out: Path) -> Path:
    df = pd.read_csv(csv_path)
    df = df.replace([np.inf, -np.inf], np.nan).dropna(subset=["log2FoldChange", "padj"])
    df["neglog10"] = -np.log10(df["padj"].clip(lower=1e-300))
    up = (df["padj"] < 0.05) & (df["log2FoldChange"] >= 1)
    down = (df["padj"] < 0.05) & (df["log2FoldChange"] <= -1)
    # Widescreen slide proportions (~16:9 content)
    fig, ax = plt.subplots(figsize=(12.6, 6.4))
    ax.scatter(
        df.loc[~up & ~down, "log2FoldChange"],
        df.loc[~up & ~down, "neglog10"],
        s=10,
        c="#c8c8c8",
        alpha=0.45,
        rasterized=True,
    )
    ax.scatter(df.loc[up, "log2FoldChange"], df.loc[up, "neglog10"], s=14, c="#c0392b", alpha=0.75, label=f"Up {int(up.sum())}")
    ax.scatter(
        df.loc[down, "log2FoldChange"],
        df.loc[down, "neglog10"],
        s=14,
        c="#2980b9",
        alpha=0.75,
        label=f"Down {int(down.sum())}",
    )
    ax.axvline(1, color="#888888", ls="--", lw=1)
    ax.axvline(-1, color="#888888", ls="--", lw=1)
    ax.axhline(-math.log10(0.05), color="#888888", ls="--", lw=1)
    ax.set_xlabel("log2 fold change", fontsize=12)
    ax.set_ylabel("-log10(padj)", fontsize=12)
    ax.set_title(title, fontsize=14)
    ax.legend(fontsize=11, loc="upper right", frameon=False)
    ax.tick_params(labelsize=10)
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=160)
    plt.close(fig)
    return out


def make_pca(csv_path: Path, out: Path) -> Path:
    df = pd.read_csv(csv_path)
    fig, ax = plt.subplots(figsize=(12.6, 6.4))
    order = ["ctrl", "4H", "24H", "72H", "7D"]
    for g in order:
        sub = df[df["group"] == g]
        if sub.empty:
            continue
        ax.scatter(
            sub["PC1"],
            sub["PC2"],
            s=90,
            c=TIME_COLORS.get(str(g), "#888888"),
            label=str(g),
            edgecolors="white",
            linewidths=0.6,
        )
    ax.set_xlabel("PC1", fontsize=12)
    ax.set_ylabel("PC2", fontsize=12)
    ax.set_title("BMP50 bulk RNA-seq PCA", fontsize=14)
    ax.legend(fontsize=11, frameon=False)
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=160)
    plt.close(fig)
    return out


def classify_trajectory(row: pd.Series) -> str:
    times = ["ctrl", "4H", "24H", "72H", "7D"]
    vals = np.array([float(row[t]) for t in times], dtype=float)
    peak = times[int(np.argmax(vals))]
    trough = times[int(np.argmin(vals))]
    # Relative to ctrl
    early_up = vals[1] > vals[0] + 0.8 or vals[2] > vals[0] + 0.8
    late_down = vals[4] < max(vals[1], vals[2]) - 0.8
    mid_down = vals[2] < vals[0] - 0.8 and vals[4] > vals[2] + 0.8
    late_up = peak == "7D" and vals[4] > vals[0] + 0.8
    if mid_down:
        return "down_then_up"
    if early_up and late_down:
        return "up_then_down"
    if late_up:
        return "late_up"
    if peak in ("4H", "24H") and trough == "7D":
        return "up_then_down"
    if trough in ("24H", "72H") and peak == "7D":
        return "down_then_up"
    return "other"


def make_lrt_trajectory_plots(mean_csv: Path, long_csv: Path) -> dict[str, Path]:
    mean = pd.read_csv(mean_csv)
    long = pd.read_csv(long_csv)
    mean["pattern"] = mean.apply(classify_trajectory, axis=1)
    genes = list(dict.fromkeys(long["gene_id"].tolist()))
    # keep only genes present in the top-LRT long table
    mean = mean[mean["gene_id"].isin(genes)].copy()
    mean["pattern"] = mean.apply(classify_trajectory, axis=1)

    outs: dict[str, Path] = {}
    # All together
    outs["lrt_all"] = _plot_gene_panel(long, genes, GEN / "lrt_all.png", "Top LRT timecourse genes (all)")

    groups = {
        "up_then_down": "Early/mid up, then down (transient induction)",
        "down_then_up": "Mid down, then recover (e.g. cell-cycle dip)",
        "late_up": "Late rise (7d peak)",
        "other": "Other trajectories",
    }
    for key, title in groups.items():
        gset = mean.loc[mean["pattern"] == key, "gene_id"].tolist()
        if not gset:
            continue
        outs[f"lrt_{key}"] = _plot_gene_panel(long, gset, GEN / f"lrt_{key}.png", title)
    # save classification
    mean[["gene_id", "pattern"]].to_csv(GEN / "lrt_gene_patterns.csv", index=False)
    return outs


def _plot_gene_panel(long: pd.DataFrame, genes: list[str], out: Path, title: str) -> Path:
    n = len(genes)
    ncols = min(4, max(1, n))
    nrows = int(math.ceil(n / ncols))
    fig, axes = plt.subplots(nrows, ncols, figsize=(12.6, max(3.2, 2.3 * nrows)), squeeze=False)
    order = ["ctrl", "4h", "24h", "72h", "7d"]
    x = np.arange(len(order))
    for i, gene in enumerate(genes):
        ax = axes[i // ncols][i % ncols]
        sub = long[long["gene_id"] == gene]
        means, sems = [], []
        for t in order:
            vals = sub.loc[sub["time"].str.lower() == t, "vst"].astype(float).values
            means.append(float(np.mean(vals)) if len(vals) else np.nan)
            sems.append(float(np.std(vals, ddof=1) / np.sqrt(len(vals))) if len(vals) > 1 else 0.0)
        ax.errorbar(x, means, yerr=sems, fmt="-o", color="#6a3d9a", lw=1.5, ms=5, capsize=2)
        ax.set_xticks(x)
        ax.set_xticklabels(order, fontsize=8)
        ax.set_title(gene, fontsize=10)
        ax.tick_params(labelsize=8)
    for j in range(n, nrows * ncols):
        axes[j // ncols][j % ncols].axis("off")
    fig.suptitle(title, fontsize=13, y=1.01)
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


def make_wide_heatmap(mean_csv: Path, out: Path) -> Path:
    df = pd.read_csv(mean_csv)
    # Use top LRT genes if available in same file (already short list in mean_vst_lrt)
    genes = df["gene_id"].tolist()
    mat = df[["ctrl", "4H", "24H", "72H", "7D"]].to_numpy(dtype=float)
    # z-score each gene across time for visualization
    z = (mat - mat.mean(axis=1, keepdims=True)) / (mat.std(axis=1, keepdims=True) + 1e-9)
    # Wide landscape: times on x, genes on y but keep height modest
    fig, ax = plt.subplots(figsize=(12.6, 4.8))
    im = ax.imshow(z, aspect="auto", cmap="RdBu_r", vmin=-2.5, vmax=2.5)
    ax.set_xticks(range(5))
    ax.set_xticklabels(["ctrl", "4h", "24h", "72h", "7d"], fontsize=11)
    ax.set_yticks(range(len(genes)))
    ax.set_yticklabels(genes, fontsize=7)
    ax.set_title("LRT timecourse genes (row z-scored VST)", fontsize=13)
    fig.colorbar(im, ax=ax, fraction=0.02, pad=0.02, label="z-score")
    fig.tight_layout()
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=160)
    plt.close(fig)
    return out


def make_program_panels(
    scores_csv: Path,
    pred_csv: Path,
    cv_csv: Path,
    metrics: dict,
    names: tuple[str, ...],
    out: Path,
    title: str,
    ncols: int,
) -> Path:
    scores = pd.read_csv(scores_csv)
    pred = pd.read_csv(pred_csv)
    cv = pd.read_csv(cv_csv)
    actual = scores.set_index("time").loc[list(TIME_ORDER), list(names)].to_numpy(dtype=float)
    pred_by_time = []
    for t in TIME_ORDER:
        slices = cv.loc[cv["true_time"] == t, "slice"].astype(str)
        block = pred[pred["slice"].astype(str).isin(slices)]
        pred_by_time.append(block[list(names)].astype(float).mean(axis=0).to_numpy())
    pred_mat = np.vstack(pred_by_time)

    n = len(names)
    nrows = int(math.ceil(n / ncols))
    fig, axes = plt.subplots(nrows, ncols, figsize=(12.6, 2.55 * nrows), sharey=True, squeeze=False)
    for i, name in enumerate(names):
        ax = axes[i // ncols][i % ncols]
        a = actual[:, i]
        p = pred_mat[:, i]
        ax.plot(TIME_ORDER, a, "-o", color="#333333", lw=1.8, label="true bulk RNA")
        ax.plot(TIME_ORDER, p, "-o", color="#6a3d9a", lw=1.8, label="from morphology")
        ax.axhline(0, color="#dddddd", lw=0.8)
        r = metrics["programs"][name]["pearson"]
        ax.set_title(f"{name.replace('_', ' ')}   r = {r:.2f}", fontsize=11)
        ax.tick_params(axis="x", rotation=30, labelsize=9)
        ax.tick_params(axis="y", labelsize=9)
    for j in range(n, nrows * ncols):
        axes[j // ncols][j % ncols].axis("off")
    axes[0][0].set_ylabel("program z-score", fontsize=10)
    handles, labels = axes[0][0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=2, fontsize=10, frameon=False, bbox_to_anchor=(0.5, -0.02))
    fig.suptitle(title, fontsize=13, y=1.01)
    fig.tight_layout(rect=(0, 0.04, 1, 0.97))
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


def prepare() -> dict[str, Path]:
    GEN.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    paths["pca"] = make_pca(RNA / "timecourse_pca.csv", GEN / "timecourse_pca.png")
    for label, fname in [
        ("4H", "4H_vs_ctrl.csv"),
        ("24H", "24H_vs_ctrl.csv"),
        ("72H", "72H_vs_ctrl.csv"),
        ("7D", "7D_vs_ctrl.csv"),
    ]:
        paths[f"volcano_{label}"] = make_volcano(
            RNA / fname, f"{label} vs ctrl", GEN / f"volcano_{label}_vs_ctrl.png"
        )
    paths["heatmap"] = make_wide_heatmap(RNA / "mean_vst_lrt.csv", GEN / "heatmap_wide.png")
    paths.update(make_lrt_trajectory_plots(RNA / "mean_vst_lrt.csv", RNA / "timecourse_top_lrt_genes.csv"))

    metrics = json.loads((MORPH / "rna_metrics.json").read_text(encoding="utf-8"))
    paths["core_3x2"] = make_program_panels(
        MORPH / "rna_program_scores.csv",
        MORPH / "predicted_programs.csv",
        MORPH / "cv_predictions.csv",
        metrics,
        CORE,
        GEN / "core_six_3x2.png",
        "Core RNA programs: true bulk (black) vs morphology-predicted (purple)",
        ncols=3,
    )
    all_names = tuple(metrics["programs"].keys())
    paths["all14"] = make_program_panels(
        MORPH / "rna_program_scores.csv",
        MORPH / "predicted_programs.csv",
        MORPH / "cv_predictions.csv",
        metrics,
        all_names,
        GEN / "all14_with_r.png",
        "All 14 RNA programs (r = Pearson recovery from morphology)",
        ncols=4,
    )
    paths["supp"] = make_program_panels(
        MORPH / "rna_program_scores.csv",
        MORPH / "predicted_programs.csv",
        MORPH / "cv_predictions.csv",
        metrics,
        SUPP,
        GEN / "supp_with_r.png",
        "Supplementary programs (immune / function / ECM)",
        ncols=2,
    )
    # cytoskeleton program alone
    paths["cyto_prog"] = make_program_panels(
        MORPH / "rna_program_scores.csv",
        MORPH / "predicted_programs.csv",
        MORPH / "cv_predictions.csv",
        metrics,
        ("cytoskeleton",),
        GEN / "cytoskeleton_program.png",
        "Cytoskeleton RNA program recovered from morphology",
        ncols=1,
    )
    return paths


def build() -> Path:
    paths = prepare()
    metrics = json.loads((MORPH / "rna_metrics.json").read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "BMP50 astrocytes: morphology → RNA programs",
        "Predict transcriptional state from imaging\n"
        "(bulk RNA-seq builds the frozen atlas; morphology drives inference)",
    )

    add_bullets(
        prs,
        "Question",
        [
            "What does BMP50 do to astrocyte transcription over time (ctrl → 4h → 24h → 72h → 7d)?",
            "Which coordinated gene programs move: BMP signaling, cell cycle, reactive state, lineage, late maturation?",
            "Imaging is cheap and rich in shape — can morphology alone recover those bulk RNA programs?",
            "If yes: how? (predict time from morph, then look up a frozen RNA atlas — not cell-matched RNA)",
            "Does skeleton morphometry itself equal the cytoskeleton gene program? Or are they related but offset in time?",
        ],
        size=15,
    )

    add_bullets(
        prs,
        "Design: morphology → RNA (RNA atlas built once)",
        [
            "Imaging pipeline: U-Net / Cellpose-style segmentation → single-cell crops → CellProfiler + skeleton morphometrics",
            "Pool cells to slice features (medians, p90s, cluster fractions, optional cell-level time helper)",
            "Stage 1 — morph → time: L2 multinomial logistic regression with nested CV (~92% slice accuracy)",
            "Stage 2 — time → RNA: predicted_RNA = P(time|morph) · frozen bulk atlas (programs + genes)",
            "Bulk RNA-seq (DESeq2 VST) only builds/validates that atlas — it is not an input when you predict on new FOVs",
            "No cell-matched RNA: imaging slices and bulk wells share time labels only",
        ],
        size=15,
    )

    add_full_image(
        prs,
        "Bulk RNA-seq PCA: samples separate by BMP50 time",
        paths["pca"],
        "Transcriptional state changes continuously along the time course — the atlas we will later mix from morphology.",
    )

    for label in ["4H", "24H", "72H", "7D"]:
        add_full_image(
            prs,
            f"Volcano: {label} vs ctrl  (padj<0.05, |LFC|≥1) — label genes by hand",
            paths[f"volcano_{label}"],
            "Red = significantly up; blue = significantly down. No gene labels (add interactively).",
        )

    add_full_image(
        prs,
        "LRT timecourse heatmap (wide): genes × time",
        paths["heatmap"],
        "Row z-scored mean VST. Shows coordinated blocks: early BMP, mid arrest, late recovery/maturation.",
    )

    add_full_image(prs, "Top LRT genes — all trajectories together", paths["lrt_all"])
    for key, title in [
        ("lrt_up_then_down", "LRT group: up then down (transient induction)"),
        ("lrt_down_then_up", "LRT group: down then up (mid dip / recovery)"),
        ("lrt_late_up", "LRT group: late up (7d peak)"),
        ("lrt_other", "LRT group: other patterns"),
    ]:
        if key in paths:
            add_full_image(prs, title, paths[key])

    add_bullets(
        prs,
        "RNA biology snapshot — what BMP50 is doing",
        [
            "Early (4h): classic BMP/SMAD pulse — Id1/Id3, Smad6/7, Nog strongly induced; oligodendrocyte-lineage TFs (e.g. Olig2) suppressed. This is “BMP is on.”",
            "Mid (24h): dominant genome-wide pattern is transient — proliferation genes (Mki67, Top2a, Ccna2…) drop (arrest), then largely recover by 7d. Gfap also rises mid-course but is not sustained forever.",
            "Mid–late (72h): reactive / injury-like and immune-adjacent axes move (Stat3, Lcn2/Serpina3n, MHC class I, complement). Homeostatic astrocyte function genes (uptake / K+ / connexin) can rise here too.",
            "Late (7d): a distinct maturation / thyroid-hormone-related set turns on (Dio2, Slco1c1, Thy1, Nnat). Very few genes are sustained in the same direction at every timepoint (Smad6 up; Dll1 down).",
            "Takeaway for imaging: morphology should be able to tell early vs mid vs late states if shape tracks these transitions — then we can recover the corresponding RNA programs.",
        ],
        size=14,
    )

    add_bullets(
        prs,
        "Bridge: how morphology maps onto RNA programs",
        [
            "A program is a gene-set score from bulk RNA (mean VST of curated genes → z-score → average by time). Example: cytoskeleton = Vim, Nes, actin/MT/Rho genes — not a skeleton length pixel count.",
            "We curated ~14 programs (BMP_Id, cell_cycle, reactive_GFAP, cytoskeleton, OPC_pdgfra, late_7d, immune axes, glutamate/ion, ECM…). Bulk RNA defines the atlas once.",
            "Morphology never reads genes. It only predicts P(time). We then mix the frozen atlas: predicted_program = Σ P(time) × atlas(time).",
            "That is why purple curves are smoother than black: uncertain time blends neighboring states — healthy generalization, not failure.",
            "Direction of the model: morph features → time → RNA programs (RNA is lookup, not input at inference).",
        ],
        size=14,
    )

    add_bullets(
        prs,
        "Morphology → time: nested CV (what we trust)",
        [
            "We do not quote a single train/test on the same slices without nesting.",
            "Honest result: nested top-35 feature selection inside each fold → ~92% slice accuracy for time.",
            "Avoid the leaky ~98% run where features were chosen using all slices before CV.",
            "Confusion matrices look impressive but still reflect CV folds of this one experiment — treat as supportive, not definitive external validation.",
            "What matters next: once time is predicted, RNA programs mix from the atlas.",
        ],
        size=15,
    )

    add_bullets(
        prs,
        "Top morph features for time — grouped",
        [
            "Cell-helper time votes (stable in 5/5 folds): mean_Pcell_ctrl, mean_Pcell_4h, mean_Pcell_24h, mean_Pcell_72h — average per-cell time probabilities pooled to the slice.",
            "GFAP / Channel_0 intensity: p90_Intensity_IntegratedIntensity_Channel_0, median_Intensity_IntegratedIntensity_Channel_0, related edge/max/std intensity — how strongly GFAP lights up.",
            "Outline / shape (Zernike & moments): p90/median AreaShape_Zernike_* (e.g. 6_0, 2_2, 5_3, 7_7, 4_4, 6_6, 8_0…), NormalizedMoment_3_0 — compact summary of process silhouette.",
            "Mass displacement / other channels: median/p90 Intensity_MassDisplacement_Channel_1/2, LowerQuartileIntensity_Channel_2, IntegratedIntensityEdge_Channel_2 — intensity geometry beyond raw GFAP sum.",
            "Population mix: frac_cluster_1, frac_cluster_2 (and sometimes 4) — which morphological clusters dominate the FOV.",
            "Skeleton / branching (selected often but not always): median_nontrunk_branch_points, p90_nontrunk_branch_points — process complexity; useful for time, but (later) not a 1:1 cytoskeleton-RNA proxy.",
        ],
        size=13,
    )

    add_full_image(
        prs,
        "Core story (3×2): true bulk vs morph-predicted programs",
        paths["core_3x2"],
        "Black = true bulk program by time; purple = morphology via P(time)·atlas. r on each panel = recovery strength.",
    )

    add_bullets(
        prs,
        "What does the r on each panel mean?",
        [
            "For each program, every imaging slice gets a predicted score from morphology (via P(time)).",
            "We compare that prediction to the true bulk program score for the slice’s real timepoint.",
            "Pearson r ≈ how well the morph-derived scores track the true time-course of that program across slices.",
            "High r (e.g. late_7d 0.96, cytoskeleton 0.94) means morphology recovers that program’s trajectory well.",
            "This is NOT “genes measured in the image” — it is atlas mixing after time prediction.",
        ],
        size=15,
    )

    add_full_image(
        prs,
        "All 14 RNA programs (with r on each panel)",
        paths["all14"],
        f"Mean recovery across programs r ≈ {metrics['mean_pearson']:.2f}. Purple smoother than black is expected.",
    )

    add_image_with_notes(
        prs,
        "Supplementary programs — with r",
        paths["supp"],
        [
            f"MHC_class_I r={metrics['programs']['MHC_class_I']['pearson']:.2f}: antigen presentation genes track mid/late reactive window.",
            f"Stat3_reactive r={metrics['programs']['Stat3_reactive']['pearson']:.2f}: IL-6–like signaling hub (not the same as GFAP structure).",
            f"glutamate_ion r={metrics['programs']['glutamate_ion']['pearson']:.2f}: uptake / K+ / gap-junction function.",
            f"ECM_remodeling r={metrics['programs']['ECM_remodeling']['pearson']:.2f}: matrix genes (e.g. Col11a2 axis).",
            "These are secondary to the BMP→arrest→reactive→late core story, but recoverable from morph.",
        ],
    )

    add_full_image(
        prs,
        "Skeleton morph vs true cytoskeleton RNA over time",
        MORPH / "cytoskeleton_vs_skeleton_corr.png",
        "True bulk cytoskeleton (purple) vs skeleton mean±SEM (black). Stars = Kruskal-Wallis across times.",
    )

    add_full_image(
        prs,
        "Cytoskeleton RNA program — recovered from morphology",
        paths["cyto_prog"],
        "Even though skeleton length ≠ cytoskeleton RNA, the cytoskeleton PROGRAM is still recovered via morph→time (r≈0.94).",
    )

    add_bullets(
        prs,
        "Interpretation: program recovery vs direct skeleton correlation",
        [
            "True cytoskeleton RNA peaks at ~24h and then falls; skeleton length / branch points keep rising into 7d.",
            "So there is no simple linear “longer skeleton ⇒ higher cytoskeleton transcription” relationship — they are out of phase.",
            "That does NOT mean morphology fails: the full morph feature set still predicts time, and therefore recovers the cytoskeleton gene program (r≈0.94).",
            "In words: morphology is a good state sensor for the transcriptional cytoskeleton package, but skeleton morphometrics alone are a late structural readout — not a substitute for that gene set.",
            "For the lab: use morph→program for RNA-state inference; use skeleton features to talk about process elaboration timing separately.",
        ],
        size=14,
    )

    add_bullets(
        prs,
        "Take-home",
        [
            "BMP50 drives a clear bulk RNA trajectory: early BMP pulse → mid cycle arrest / reactive shift → late maturation.",
            "We summarize that biology as curated RNA programs (gene-set scores), not as morph features.",
            "Astrocyte morphology (U-Net/CellProfiler pipeline → logistic time model) predicts time at ~92% (nested CV) and recovers 14 programs at mean r≈0.90.",
            "Purple predicted curves are smoother because they mix times — expected for this atlas design.",
            "Skeleton morph and cytoskeleton transcription both respond to BMP but peak at different times — related stories, not the same variable.",
            "Practical next step: new FOVs → morph CSV only → predict_rna_from_morph.py with the frozen bundle (no new RNA required).",
        ],
        size=14,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(f"Wrote {build()}")
