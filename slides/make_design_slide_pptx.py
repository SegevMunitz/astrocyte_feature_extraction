#!/usr/bin/env python3
"""Editable Design slide: morph → RNA as native PowerPoint shapes."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

OUT = Path(__file__).resolve().parents[1] / "slides" / "design_morph_to_RNA_editable.pptx"

# Colors
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x55, 0x55, 0x55)
BLUE = RGBColor(0xEE, 0xF2, 0xF7)
ORANGE = RGBColor(0xF7, 0xEB, 0xE0)
PURPLE = RGBColor(0xEF, 0xE6, 0xF5)
GREEN = RGBColor(0xE8, 0xF0, 0xE8)
CREAM = RGBColor(0xF5, 0xF0, 0xE6)
LINE = RGBColor(0x33, 0x33, 0x33)


def _set_fill(shape, rgb: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb


def _set_line(shape, rgb: RGBColor = LINE, width_pt: float = 1.25) -> None:
    shape.line.color.rgb = rgb
    shape.line.width = Pt(width_pt)


def rounded_box(slide, left, top, width, height, text: str, fill: RGBColor, size: int = 12, bold: bool = False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    _set_fill(shape, fill)
    _set_line(shape)
    # milder corner radius
    try:
        adj = shape.adjustments
        adj[0] = 0.15
    except Exception:
        pass
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].clear()
    # vertical center via anchor
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")  # type: ignore[attr-defined]
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for i, line in enumerate(text.split("\n")):
        para = p if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = INK
        run.font.name = "Calibri"
        para.space_before = Pt(0)
        para.space_after = Pt(0)
    return shape


def label(slide, left, top, width, height, text: str, size: int = 14, bold: bool = True, color=INK):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def h_arrow(slide, x1, y_center, x2):
    """Straight horizontal connector with arrow head."""
    # slight vertical thickness via connector
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        x1,
        y_center,
        x2,
        y_center,
    )
    connector.line.color.rgb = LINE
    connector.line.width = Pt(1.75)
    # arrow end
    ln = connector.line._ln
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return connector


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Title
    label(slide, Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.45), "Design: morphology  →  RNA", size=26, bold=True)
    label(
        slide,
        Inches(0.4),
        Inches(0.65),
        Inches(12.5),
        Inches(0.35),
        "Imaging predicts time · frozen bulk atlas supplies programs · no cell-matched RNA",
        size=13,
        bold=False,
        color=MUTED,
    )

    # --- A. Imaging ---
    label(slide, Inches(0.4), Inches(1.15), Inches(12), Inches(0.3), "A. Imaging → slice features", size=14, bold=True)
    yA = Inches(1.5)
    hA = Inches(0.85)
    wA = Inches(2.7)
    gap = Inches(0.35)
    x0 = Inches(0.4)
    a_labels = [
        "Cellpose\nsegmentation",
        "Single-cell\ncrops",
        "CellProfiler +\nskeleton metrics",
        "Pool to slice\n(median / p90 / clusters)",
    ]
    xs = []
    for i, text in enumerate(a_labels):
        x = x0 + i * (wA + gap)
        xs.append(x)
        rounded_box(slide, x, yA, wA, hA, text, BLUE, size=11)
    for i in range(3):
        h_arrow(slide, xs[i] + wA, yA + hA / 2, xs[i + 1])

    # --- B. Two-stage ---
    label(
        slide,
        Inches(0.4),
        Inches(2.55),
        Inches(12),
        Inches(0.3),
        "B. Two-stage model (inference = morph only)",
        size=14,
        bold=True,
    )
    yB = Inches(2.95)
    hB = Inches(1.05)
    boxes_B = [
        (Inches(0.4), Inches(2.6), "Slice morphology\nfeatures", GREEN, 12),
        (Inches(3.55), Inches(3.5), "Stage 1\nMorph → P(time)\nL2 logistic · nested CV ~92%", ORANGE, 11),
        (Inches(7.55), Inches(2.35), "Mix\nP(time) · atlas", PURPLE, 12),
        (Inches(10.4), Inches(2.5), "Predicted RNA\nprograms + genes", GREEN, 11),
    ]
    xb = []
    for x, w, text, fill, sz in boxes_B:
        rounded_box(slide, x, yB, w, hB, text, fill, size=sz, bold=True)
        xb.append((x, w))
    for i in range(3):
        x1 = xb[i][0] + xb[i][1]
        x2 = xb[i + 1][0]
        h_arrow(slide, x1, yB + hB / 2, x2)

    # --- C. Atlas ---
    label(
        slide,
        Inches(0.4),
        Inches(4.25),
        Inches(12),
        Inches(0.3),
        "C. Frozen RNA atlas (built once from bulk RNA-seq)",
        size=14,
        bold=True,
    )
    yC = Inches(4.65)
    hC = Inches(1.15)
    c_boxes = [
        (Inches(0.4), Inches(3.2), "Bulk wells\nDESeq2 → VST", CREAM, 12),
        (Inches(4.15), Inches(4.0), "Gene-set programs\n(+ panel genes)\nz-scored by time", CREAM, 11),
        (Inches(8.7), Inches(4.2), "Atlas lookup\n5 times × K programs\nNot an input on new FOVs", CREAM, 11),
    ]
    xc = []
    for x, w, text, fill, sz in c_boxes:
        rounded_box(slide, x, yC, w, hC, text, fill, size=sz)
        xc.append((x, w))
    for i in range(2):
        h_arrow(slide, xc[i][0] + xc[i][1], yC + hC / 2, xc[i + 1][0])

    # Dashed note linking atlas → mix (textbox as callout; user can redraw connector)
    note = slide.shapes.add_textbox(Inches(8.9), Inches(4.05), Inches(2.8), Inches(0.3))
    p = note.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "↑ feeds Stage 2 mix"
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x6A, 0x3D, 0x9A)
    run.font.name = "Calibri"

    # Footer
    label(
        slide,
        Inches(0.4),
        Inches(6.15),
        Inches(12.5),
        Inches(0.55),
        "Constraints: imaging slices & bulk wells share time labels only  ·  no cell-matched transcriptomes  ·  new FOVs → morph CSV only",
        size=12,
        bold=False,
        color=MUTED,
    )
    label(
        slide,
        Inches(0.4),
        Inches(6.7),
        Inches(12.5),
        Inches(0.4),
        "Tip: all boxes/text are editable shapes — move, recolor, or rewrite freely.",
        size=10,
        bold=False,
        color=RGBColor(0x99, 0x99, 0x99),
    )

    # Optional second blank slide? No — one slide deliverable
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
